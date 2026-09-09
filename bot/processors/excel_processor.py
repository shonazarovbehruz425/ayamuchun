import asyncio
import os
import csv
import logging
from openpyxl import load_workbook, Workbook

from .base import BaseProcessor, FileMetadata

logger = logging.getLogger(__name__)

class ExcelProcessor(BaseProcessor):
    """
    Excel hujjatlarni qayta ishlash uchun sinf.
    """
    SUPPORTED_EXTENSIONS = ['.xlsx', '.xls', '.xlsm']

    async def extract_text(self, file_path: str) -> str:
        """Excel dan matnni ajratib olish (barcha hujayralardan)."""
        def _extract():
            wb = load_workbook(file_path, data_only=True)
            text = []
            for sheet_name in wb.sheetnames:
                sheet = wb[sheet_name]
                text.append(f"--- Sheet: {sheet_name} ---")
                for row in sheet.iter_rows(values_only=True):
                    row_data = [str(cell) for cell in row if cell is not None]
                    if row_data:
                        text.append("\t".join(row_data))
            return "\n".join(text)
            
        try:
            return await asyncio.to_thread(_extract)
        except Exception as e:
            logger.error(f"Xato: Excel dan matnni ajratishda xatolik - {e}")
            raise

    async def get_metadata(self, file_path: str) -> FileMetadata:
        """Excel metama'lumotlarini olish."""
        def _get_meta():
            wb = load_workbook(file_path, read_only=True)
            size = os.path.getsize(file_path)
            
            sheet_count = len(wb.sheetnames)
            
            # Faqat faol varaq qatorlari hisobi yoki hamma
            active_sheet = wb.active
            row_count = active_sheet.max_row
            col_count = active_sheet.max_column
            
            return FileMetadata(
                file_name=os.path.basename(file_path),
                file_type="Excel Spreadsheet",
                file_size=size,
                page_count=sheet_count,  # Varoqlar sonini page_count da ko'rsatamiz
                has_tables=True,
                extra={"row_count": row_count, "column_count": col_count}
            )
            
        try:
            return await asyncio.to_thread(_get_meta)
        except Exception as e:
            logger.error(f"Xato: Excel metama'lumotlarini olishda xatolik - {e}")
            raise

    async def read_data(self, file_path: str, sheet_name: str = None) -> list[list]:
        """Excel hujjatidan ma'lumotlarni o'qish."""
        def _read():
            wb = load_workbook(file_path, data_only=True)
            sheet = wb[sheet_name] if sheet_name and sheet_name in wb.sheetnames else wb.active
            
            data = []
            for row in sheet.iter_rows(values_only=True):
                data.append(list(row))
            return data
            
        try:
            return await asyncio.to_thread(_read)
        except Exception as e:
            logger.error(f"Xato: Excel ma'lumotlarini o'qishda xatolik - {e}")
            raise

    async def create_spreadsheet(self, data: list[list], output_path: str, headers: list[str] = None, sheet_name: str = 'Sheet1') -> str:
        """Yangi Excel hujjatini yaratish."""
        def _create():
            wb = Workbook()
            sheet = wb.active
            sheet.title = sheet_name
            
            start_row = 1
            if headers:
                sheet.append(headers)
                start_row = 2
                
            for row in data:
                sheet.append(row)
                
            wb.save(output_path)
            return output_path
            
        try:
            return await asyncio.to_thread(_create)
        except Exception as e:
            logger.error(f"Xato: Excel hujjati yaratishda xatolik - {e}")
            raise

    async def calculate_stats(self, file_path: str, column: int = None) -> dict:
        """Ustunlar bo'yicha statistika hisoblash (barcha raqamli ustunlar yoki ko'rsatilgan ustun)."""
        def _stats():
            ext = os.path.splitext(file_path)[1].lower()
            data_rows = []
            
            if ext == '.csv':
                for enc in ('utf-8-sig', 'utf-8', 'cp1251', 'latin-1'):
                    try:
                        with open(file_path, 'r', encoding=enc, errors='replace') as f:
                            r = csv.reader(f)
                            for row in r:
                                data_rows.append(row)
                        break
                    except Exception:
                        continue
            else:
                wb = load_workbook(file_path, data_only=True)
                sheet = wb.active
                for row in sheet.iter_rows(values_only=True):
                    data_rows.append(list(row))

            if not data_rows:
                return {}

            headers = [str(c or f"Ustun {i+1}") for i, c in enumerate(data_rows[0])]
            rows_data = data_rows[1:] if len(data_rows) > 1 else []

            columns_stats = {}

            if column is not None:
                target_indices = [column - 1]
            else:
                target_indices = list(range(len(headers)))

            for idx in target_indices:
                if idx < 0 or idx >= len(headers):
                    continue
                col_name = headers[idx]
                num_values = []
                for row in rows_data:
                    if idx < len(row):
                        val = row[idx]
                        if isinstance(val, (int, float)):
                            num_values.append(float(val))
                        elif isinstance(val, str):
                            val_str = val.strip().replace(',', '.')
                            try:
                                num_values.append(float(val_str))
                            except ValueError:
                                pass
                
                if num_values:
                    columns_stats[col_name] = {
                        "Eng kichik (min)": min(num_values),
                        "Eng katta (max)": max(num_values),
                        "O'rtacha (avg)": round(sum(num_values) / len(num_values), 2),
                        "Jami yig'indi (sum)": round(sum(num_values), 2),
                        "Raqamlar soni": len(num_values)
                    }

            if not columns_stats and column is not None:
                return {"min": 0, "max": 0, "avg": 0, "count": 0}

            return columns_stats

        try:
            return await asyncio.to_thread(_stats)
        except Exception as e:
            logger.error(f"Xato: Statistika hisoblashda xatolik - {e}")
            raise

    async def create_grade_table(self, students: list[str], subjects: list[str], output_path: str) -> str:
        """Baholash uchun shablon yaratish."""
        def _create_grade():
            wb = Workbook()
            sheet = wb.active
            sheet.title = "Baholar"
            
            headers = ["O'quvchi"] + subjects
            sheet.append(headers)
            
            for student in students:
                row = [student] + [""] * len(subjects)
                sheet.append(row)
                
            wb.save(output_path)
            return output_path
            
        try:
            return await asyncio.to_thread(_create_grade)
        except Exception as e:
            logger.error(f"Xato: Baholar jadvalini yaratishda xatolik - {e}")
            raise

    async def to_csv(self, file_path: str, output_path: str) -> str:
        """Excel dan CSV formatiga o'tkazish."""
        def _convert():
            wb = load_workbook(file_path, data_only=True)
            sheet = wb.active
            
            with open(output_path, mode='w', newline='', encoding='utf-8') as f:
                writer = csv.writer(f)
                for row in sheet.iter_rows(values_only=True):
                    writer.writerow(["" if cell is None else str(cell) for cell in row])
            return output_path
            
        try:
            return await asyncio.to_thread(_convert)
        except Exception as e:
            logger.error(f"Xato: CSV formatiga o'tkazishda xatolik - {e}")
            raise

    async def convert_to_word_table(
        self,
        file_path: str,
        output_path: str,
        orientation_mode: str = 'auto',
        table_style: str = 'modern_blue'
    ) -> dict:
        """
        Excel (.xlsx, .xls) yoki CSV jadvallarini tahlil qilib,
        orientatsiyani (Albom/Kitob) avtomatik moslab, ustunlar enini hisoblab,
        zamonaviy va chiroyli Word (.docx) hujjatiga o'tkazish.
        """
        def _convert_worker():
            from docx import Document
            from docx.shared import Inches, Pt, RGBColor
            from docx.enum.section import WD_ORIENT
            from docx.enum.text import WD_ALIGN_PARAGRAPH
            from docx.enum.table import WD_TABLE_ALIGNMENT
            from docx.oxml import parse_xml
            from docx.oxml.ns import nsdecls

            ext = os.path.splitext(file_path)[1].lower()
            sheets_data = []  # list of (sheet_name, list_of_rows)

            # 1. Ma'lumotlarni o'qish (.xlsx, .xls, .csv)
            if ext == '.csv':
                rows = []
                for enc in ('utf-8-sig', 'utf-8', 'cp1251', 'latin-1'):
                    try:
                        with open(file_path, 'r', encoding=enc, errors='replace') as f:
                            r = csv.reader(f)
                            for row in r:
                                rows.append(row)
                        break
                    except Exception:
                        continue
                sheets_data.append(("Jadval", rows))
            elif ext == '.xls':
                try:
                    import xlrd
                    book = xlrd.open_workbook(file_path)
                    for sh_name in book.sheet_names():
                        sh = book.sheet_by_name(sh_name)
                        sheet_rows = []
                        for r_idx in range(sh.nrows):
                            sheet_rows.append(sh.row_values(r_idx))
                        sheets_data.append((sh_name, sheet_rows))
                except Exception as e:
                    logger.warning(f"xlrd orqali o'qishda ogohlantirish: {e}, openpyxl sinab ko'riladi")
                    wb = load_workbook(file_path, data_only=True)
                    for sh_name in wb.sheetnames:
                        sheet = wb[sh_name]
                        sheets_data.append((sh_name, list(sheet.iter_rows(values_only=True))))
            else:
                wb = load_workbook(file_path, data_only=True)
                for sh_name in wb.sheetnames:
                    sheet = wb[sh_name]
                    sheets_data.append((sh_name, list(sheet.iter_rows(values_only=True))))

            # 2. Uslub parametrlarini tanlash
            styles_map = {
                'modern_blue': {
                    'hdr_bg': '1E40AF',
                    'hdr_text': RGBColor(255, 255, 255),
                    'zebra_bg': 'F0F7FF',
                    'border_outer': '93C5FD',
                    'border_inner': 'E2E8F0',
                    'title_color': RGBColor(30, 64, 175)
                },
                'emerald': {
                    'hdr_bg': '065F46',
                    'hdr_text': RGBColor(255, 255, 255),
                    'zebra_bg': 'ECFDF5',
                    'border_outer': 'A7F3D0',
                    'border_inner': 'E2E8F0',
                    'title_color': RGBColor(6, 95, 70)
                },
                'classic_slate': {
                    'hdr_bg': '334155',
                    'hdr_text': RGBColor(255, 255, 255),
                    'zebra_bg': 'F8FAFC',
                    'border_outer': 'CBD5E1',
                    'border_inner': 'E2E8F0',
                    'title_color': RGBColor(30, 41, 59)
                }
            }
            cur_style = styles_map.get(table_style, styles_map['modern_blue'])

            doc = Document()
            total_processed_rows = 0
            max_cols_overall = 0
            applied_orientations = []

            first_sheet_handled = False

            for sh_idx, (sh_name, raw_rows) in enumerate(sheets_data):
                # Bo'sh qatorlarni filtrlash
                cleaned_rows = []
                for r in raw_rows:
                    if r is None:
                        continue
                    row_vals = []
                    for c in r:
                        if c is None:
                            row_vals.append("")
                        elif isinstance(c, float) and c.is_integer():
                            row_vals.append(str(int(c)))
                        else:
                            row_vals.append(str(c).strip())
                    # Qatorda birorta bo'lsa ham mazmun bormi?
                    if any(v != "" for v in row_vals):
                        cleaned_rows.append(row_vals)

                if not cleaned_rows:
                    continue

                raw_max_cols = max(len(r) for r in cleaned_rows)
                # Butunlay bo'sh ustunlarni olib tashlash
                valid_col_indices = []
                for c_idx in range(raw_max_cols):
                    has_val = False
                    for r in cleaned_rows:
                        if c_idx < len(r) and r[c_idx] != "":
                            has_val = True
                            break
                    if has_val:
                        valid_col_indices.append(c_idx)

                if not valid_col_indices:
                    continue

                col_count = len(valid_col_indices)
                max_cols_overall = max(max_cols_overall, col_count)
                total_processed_rows += len(cleaned_rows)

                # 3. Orientatsiyani avtomatik aniqlash
                # Matnlarning umumiy sig'imi va ustunlar soniga qarab
                col_max_lens = []
                for col_i in valid_col_indices:
                    m_len = max(len(r[col_i]) if col_i < len(r) else 0 for r in cleaned_rows)
                    col_max_lens.append(max(m_len, 3))
                total_char_width = sum(col_max_lens)

                if orientation_mode == 'landscape':
                    is_landscape = True
                elif orientation_mode == 'portrait':
                    is_landscape = False
                else:  # auto
                    # 7 yoki undan ortiq ustun bo'lsa yoki umumiy eni 70 belgidan oshsa -> Albom
                    if col_count >= 7 or total_char_width >= 70:
                        is_landscape = True
                    else:
                        is_landscape = False

                applied_orientations.append("Albom (Landscape)" if is_landscape else "Kitob (Portrait)")

                # Section sozlash
                if not first_sheet_handled:
                    section = doc.sections[0]
                    first_sheet_handled = True
                else:
                    section = doc.add_section()

                section.top_margin = Inches(0.5)
                section.bottom_margin = Inches(0.5)
                section.left_margin = Inches(0.5)
                section.right_margin = Inches(0.5)

                if is_landscape:
                    section.orientation = WD_ORIENT.LANDSCAPE
                    section.page_width = Inches(11.69)
                    section.page_height = Inches(8.27)
                    avail_printable_width = 10.69
                else:
                    section.orientation = WD_ORIENT.PORTRAIT
                    section.page_width = Inches(8.27)
                    section.page_height = Inches(11.69)
                    avail_printable_width = 7.27

                # Varaq sarlavhasi
                p_title = doc.add_paragraph()
                p_title.paragraph_format.space_before = Pt(4)
                p_title.paragraph_format.space_after = Pt(6)
                r_title = p_title.add_run(f"📊 {sh_name}")
                r_title.bold = True
                r_title.font.name = "Calibri"
                r_title.font.size = Pt(13)
                r_title.font.color.rgb = cur_style['title_color']

                r_badge = p_title.add_run(f"  •  {len(cleaned_rows)} qator, {col_count} ustun ({'Albom' if is_landscape else 'Kitob'})")
                r_badge.font.name = "Calibri"
                r_badge.font.size = Pt(9.5)
                r_badge.font.color.rgb = RGBColor(100, 116, 139)

                # Ustunlar enini mutanosib hisoblash
                weights = [max(l, 4) for l in col_max_lens]
                sum_w = sum(weights)
                col_widths_inches = [(w / sum_w) * avail_printable_width for w in weights]

                # Minimal enni kafolatlash (katakcha 0.45 dyuymdan kam bo'lib siqilib ketmasin)
                min_w = 0.45
                for i in range(len(col_widths_inches)):
                    if col_widths_inches[i] < min_w:
                        col_widths_inches[i] = min_w
                # Qayta mutanosiblashtirish
                cur_sum = sum(col_widths_inches)
                col_widths = [Inches((w / cur_sum) * avail_printable_width) for w in col_widths_inches]

                # Shrift hajmini ustunlar soniga qarab tanlash
                if col_count <= 5:
                    body_font_size = Pt(9.5)
                    hdr_font_size = Pt(10)
                    cell_pad_v = "120"
                    cell_pad_h = "140"
                elif col_count <= 8:
                    body_font_size = Pt(8.5)
                    hdr_font_size = Pt(9)
                    cell_pad_v = "100"
                    cell_pad_h = "120"
                elif col_count <= 12:
                    body_font_size = Pt(7.5)
                    hdr_font_size = Pt(8)
                    cell_pad_v = "80"
                    cell_pad_h = "90"
                else:
                    body_font_size = Pt(7)
                    hdr_font_size = Pt(7.5)
                    cell_pad_v = "60"
                    cell_pad_h = "70"

                # Jadval yaratish
                table = doc.add_table(rows=0, cols=col_count)
                table.alignment = WD_TABLE_ALIGNMENT.CENTER

                # Jadval umumiy chegaralari va katakcha oraliqlari
                tblPr = table._tbl.tblPr
                borders_xml = parse_xml(
                    f'<w:tblBorders {nsdecls("w")}>'
                    f'  <w:top w:val="single" w:sz="6" w:space="0" w:color="{cur_style["border_outer"]}"/>'
                    f'  <w:left w:val="single" w:sz="6" w:space="0" w:color="{cur_style["border_outer"]}"/>'
                    f'  <w:bottom w:val="single" w:sz="6" w:space="0" w:color="{cur_style["border_outer"]}"/>'
                    f'  <w:right w:val="single" w:sz="6" w:space="0" w:color="{cur_style["border_outer"]}"/>'
                    f'  <w:insideH w:val="single" w:sz="4" w:space="0" w:color="{cur_style["border_inner"]}"/>'
                    f'  <w:insideV w:val="single" w:sz="4" w:space="0" w:color="{cur_style["border_inner"]}"/>'
                    f'</w:tblBorders>'
                )
                tblPr.append(borders_xml)

                cell_mar_xml = parse_xml(
                    f'<w:tblCellMar {nsdecls("w")}>'
                    f'  <w:top w:w="{cell_pad_v}" w:type="dxa"/>'
                    f'  <w:bottom w:w="{cell_pad_v}" w:type="dxa"/>'
                    f'  <w:left w:w="{cell_pad_h}" w:type="dxa"/>'
                    f'  <w:right w:w="{cell_pad_h}" w:type="dxa"/>'
                    f'</w:tblCellMar>'
                )
                tblPr.append(cell_mar_xml)

                # Qatorlarni to'ldirish
                for r_idx, orig_row in enumerate(cleaned_rows):
                    row_obj = table.add_row()
                    trPr = row_obj._tr.get_or_add_trPr()
                    trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))

                    is_header = (r_idx == 0)
                    if is_header:
                        # Har sahifa boshida sarlavha takrorlansin
                        trPr.append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))

                    bg_color = cur_style['hdr_bg'] if is_header else (cur_style['zebra_bg'] if r_idx % 2 == 1 else "FFFFFF")

                    for target_col_idx, orig_col_idx in enumerate(valid_col_indices):
                        cell = row_obj.cells[target_col_idx]
                        cell.width = col_widths[target_col_idx]

                        raw_v = orig_row[orig_col_idx] if orig_col_idx < len(orig_row) else ""
                        cell.text = raw_v

                        # Shading
                        tcPr = cell._tc.get_or_add_tcPr()
                        tcPr.append(parse_xml(f'<w:shd {nsdecls("w")} w:fill="{bg_color}"/>'))

                        p = cell.paragraphs[0]
                        p.paragraph_format.space_before = Pt(1.5)
                        p.paragraph_format.space_after = Pt(1.5)

                        # Hizalanish (Alignment)
                        if is_header:
                            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                        else:
                            # Raqamli qiymatlarni o'ng tomonga, qolganini chap/markazga
                            clean_v = raw_v.replace(' ', '').replace(',', '.')
                            try:
                                float(clean_v)
                                p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                            except ValueError:
                                if len(raw_v) <= 4 and raw_v.isalnum():
                                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                                else:
                                    p.alignment = WD_ALIGN_PARAGRAPH.LEFT

                        # Shrift parametrlari
                        for run in p.runs:
                            run.font.name = "Calibri"
                            if is_header:
                                run.bold = True
                                run.font.size = hdr_font_size
                                run.font.color.rgb = cur_style['hdr_text']
                            else:
                                run.font.size = body_font_size
                                run.font.color.rgb = RGBColor(30, 41, 59)

                # Jadvaldan keyin biroz bo'shliq
                p_space = doc.add_paragraph()
                p_space.paragraph_format.space_after = Pt(10)

            os.makedirs(os.path.dirname(output_path) or '.', exist_ok=True)
            doc.save(output_path)

            primary_orient = "Albom (Landscape)" if "Albom (Landscape)" in applied_orientations else "Kitob (Portrait)"
            return {
                "sheets_count": len(applied_orientations),
                "total_rows": total_processed_rows,
                "total_cols": max_cols_overall,
                "primary_orientation": primary_orient,
                "output_path": output_path
            }

        try:
            return await asyncio.to_thread(_convert_worker)
        except Exception as e:
            logger.error(f"Xato: Excel jadvallarini Word ga aylantirishda xatolik - {e}")
            raise

